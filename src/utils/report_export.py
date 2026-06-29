from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.data.repositories.report_repository import (
    FinanceAssociatePerformanceRow,
    InvoiceReportRow,
)


def _format_date(
    value: date | None,
) -> str:
    if value is None:
        return ""

    return value.isoformat()


def _format_decimal(
    value: Decimal | None,
) -> float | str:
    if value is None:
        return ""

    return float(
        value,
    )


def _autosize_columns(
    worksheet,
    *,
    column_count: int,
) -> None:
    for index in range(
        1,
        column_count + 1,
    ):
        letter = get_column_letter(
            index,
        )
        max_length = 0

        for cell in worksheet[letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value,
                        ),
                    ),
                )

        worksheet.column_dimensions[letter].width = min(
            max(
                max_length + 2,
                12,
            ),
            40,
        )


def build_invoice_report_workbook(
    rows: list[InvoiceReportRow],
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice Processing"

    headers = [
        "Invoice Number",
        "Invoice Date",
        "Due Date",
        "Invoice Status",
        "Validation Outcome",
        "Total Amount",
        "Tax Amount",
        "Currency",
        "Vendor Name",
        "Vendor GSTIN",
        "Vendor Email",
        "Company Name",
        "PO Number",
        "PO Date",
        "PO Status",
        "Assigned Finance Associate",
        "Assigned Finance Manager",
        "Resolution Type",
        "Validation Issue Count",
        "Highest Issue Severity",
        "Workflow Status",
        "Approved By",
        "Rejected By",
        "Escalated By",
    ]

    worksheet.append(
        headers,
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
        )

    for row in rows:
        worksheet.append(
            [
                row.invoice_number or "",
                _format_date(
                    row.invoice_date,
                ),
                _format_date(
                    row.due_date,
                ),
                row.invoice_status or "",
                row.validation_outcome or "",
                _format_decimal(
                    row.total_amount,
                ),
                _format_decimal(
                    row.tax_amount,
                ),
                row.currency or "",
                row.vendor_name or "",
                row.vendor_gstin or "",
                row.vendor_email or "",
                row.company_name or "",
                row.po_number or "",
                _format_date(
                    row.po_date,
                ),
                row.po_status or "",
                row.assigned_finance_associate or "",
                row.assigned_finance_manager or "",
                row.resolution_type or "",
                row.validation_issue_count,
                row.highest_issue_severity or "",
                row.workflow_status or "",
                row.approved_by or "",
                row.rejected_by or "",
                row.escalated_by or "",
            ],
        )

    _autosize_columns(
        worksheet,
        column_count=len(
            headers,
        ),
    )

    buffer = BytesIO()
    workbook.save(
        buffer,
    )

    return buffer.getvalue()


def build_associate_performance_workbook(
    rows: list[FinanceAssociatePerformanceRow],
    *,
    approval_rates: list[float],
    rejection_rates: list[float],
    daywise: bool = False,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Associate Performance"

    headers = [
        "Associate Name",
    ]

    if daywise:
        headers.append(
            "Report Date",
        )

    headers.extend(
        [
            "Total Assigned Invoices",
            "Approved",
            "Rejected",
            "Needs Review",
            "Ready for Approval",
            "Ready to Pay",
            "Overdue",
            "Escalated",
            "Resolved Count",
            "Recovered Count",
            "Approval Rate (%)",
            "Rejection Rate (%)",
        ],
    )

    worksheet.append(
        headers,
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
        )

    for index, row in enumerate(
        rows,
    ):
        row_values = [
            row.associate_name,
        ]

        if daywise:
            row_values.append(
                _format_date(
                    row.report_date,
                ),
            )

        row_values.extend(
            [
                row.total_assigned,
                row.approved,
                row.rejected,
                row.needs_review,
                row.ready_for_approval,
                row.ready_to_pay,
                row.overdue,
                row.escalated,
                row.resolved_count,
                row.recovered_count,
                round(
                    approval_rates[index],
                    1,
                ),
                round(
                    rejection_rates[index],
                    1,
                ),
            ],
        )

        worksheet.append(
            row_values,
        )

    _autosize_columns(
        worksheet,
        column_count=len(
            headers,
        ),
    )

    buffer = BytesIO()
    workbook.save(
        buffer,
    )

    return buffer.getvalue()
